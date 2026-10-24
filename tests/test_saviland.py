# -*- coding: utf-8 -*-
"""测试套件：把开发过程中**真实踩过的坑**全部钉成回归测试。

设计原则
--------
1. **不联网**。所有用例都能离线通过 —— 否则评审者在他机器上跑出红叉，
   而原因只是他家网络不通。
2. **隔离数据**。测试用临时目录里的 SQLite，绝不碰 `data/saviland.db`。
3. **钉住 bug，而不只是钉住功能**。下面带 `# 回归：` 注释的用例，每一个都对应
   一个真实发生过的错误（写错了口径、写错了缓存、写错了遮罩），
   这类测试才是真正会拦住下一次事故的东西。

运行：
    <底座venv>\\Scripts\\python.exe tests\\run_all.py
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import unittest
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from PIL import Image  # noqa: E402

from saviland import config, db, llm, seed  # noqa: E402
from saviland.domain import ads as ads_domain  # noqa: E402
from saviland.domain import catalog as catalog_domain  # noqa: E402
from saviland.domain import inventory as inv_domain  # noqa: E402
from saviland.domain import orders as orders_domain  # noqa: E402
from saviland.pipelines import compliance  # noqa: E402
from saviland.render import hero as hero_render  # noqa: E402


# ==========================================================================
class TempWorkspace(unittest.TestCase):
    """把数据库与产物目录重定向到临时目录的基类。"""

    tmp: Path
    _orig: dict

    @classmethod
    def setUpClass(cls) -> None:
        # ★ 两个环境约束导致的写法，都是实测踩出来的：
        #   1) 临时目录必须建在**项目内部** —— 系统临时目录在受限沙箱下不可写
        #      （WinError 5 拒绝访问）。
        #   2) 不能用 `tempfile.mkdtemp()` —— 它以 0o700 建目录，在 Windows 沙箱下
        #      这个权限位映射成的 ACL 会让**该目录内无法再建子目录**（同样报 WinError 5），
        #      而测试正需要往里面建 ads/ outputs/ 等子目录。用普通 mkdir 建即可（走继承的 ACL）。
        import uuid

        tmp_root = ROOT / "data" / "_test_tmp"
        tmp_root.mkdir(parents=True, exist_ok=True)
        cls.tmp = tmp_root / f"run_{uuid.uuid4().hex[:8]}"
        cls.tmp.mkdir(parents=True, exist_ok=True)
        cls._orig = {
            "DB_PATH": config.DB_PATH,
            "ADS_DIR": config.ADS_DIR,
            "OUTPUTS_DIR": config.OUTPUTS_DIR,
            "ASSETS_DIR": config.ASSETS_DIR,
        }
        config.DB_PATH = cls.tmp / "test.db"
        config.ADS_DIR = cls.tmp / "ads"
        config.OUTPUTS_DIR = cls.tmp / "outputs"
        config.ASSETS_DIR = cls.tmp / "assets"
        for d in (config.ADS_DIR, config.OUTPUTS_DIR, config.ASSETS_DIR):
            d.mkdir(parents=True, exist_ok=True)
        db.init_db()

    @classmethod
    def tearDownClass(cls) -> None:
        for key, value in cls._orig.items():
            setattr(config, key, value)
        shutil.rmtree(cls.tmp, ignore_errors=True)


# ==========================================================================
class TestCompliance(unittest.TestCase):
    """合规规则：既不能漏拦，也不能误杀。"""

    def test_medical_claim_is_caught(self) -> None:
        vs = compliance.check("This lamp cures fungal infections fast.")
        self.assertTrue(any("医疗功效" in v.category for v in vs),
                        f"医疗宣称应当被拦下，实际命中：{[v.term for v in vs]}")

    def test_absolute_claim_is_caught(self) -> None:
        self.assertTrue(compliance.check("The best nail lamp ever"))
        self.assertTrue(compliance.check("Guaranteed salon results"))
        self.assertTrue(compliance.check("World's best UV lamp"))

    # 回归：第一版把所有 `cure` 都当医疗宣称替换，于是"60-second cure time"
    #（美甲行业的"固化时间"）被改成 "60-second helps improve the look of time"。
    # 违禁词没了，文案也毁了。这里钉住"行业合法用法不得误杀"。
    def test_cure_in_nail_industry_sense_is_not_flagged(self) -> None:
        for text in (
            "60-second cure time for a full set",
            "UV/LED curing lamp with low heat mode",
            "Cures gel polish evenly in one pass",
            "Fast curing, no missed spots",
        ):
            with self.subTest(text=text):
                self.assertEqual(compliance.check(text), [],
                                 f"不该误判：{text}")

    def test_fda_claim_is_caught(self) -> None:
        vs = compliance.check("FDA approved nail lamp")
        self.assertTrue(any("认证" in v.category for v in vs))

    def test_competitor_bashing_is_caught(self) -> None:
        self.assertTrue(compliance.check("Better than SUNUV lamps"))

    def test_sanitize_replaces_and_keeps_text_alive(self) -> None:
        cleaned, vs = compliance.sanitize("The best lamp, guaranteed!")
        self.assertNotIn("best", cleaned.lower())
        self.assertNotIn("guaranteed", cleaned.lower())
        self.assertTrue(cleaned.strip(), "替换后不应该变成空字符串")
        self.assertEqual(len(vs), 2)

    def test_promotional_term_is_replaced_not_deleted(self) -> None:
        # 回归：第一版把促销词替换成空串，导致 CTA 字段整个变空、主图缺一块内容
        cleaned, _ = compliance.sanitize("Buy Now")
        self.assertTrue(cleaned.strip(), "促销词替换后不应变成空串")

    def test_sanitize_copy_covers_keywords(self) -> None:
        copy = {
            "title": "Great lamp",
            "bullets": ["best value", "nice"],
            "keywords": ["nail lamp", "best nail lamp"],
            "cta": "Buy Now",
        }
        cleaned, records = compliance.sanitize_copy(copy)
        self.assertNotIn("best", " ".join(cleaned["keywords"]).lower())
        self.assertTrue(any(r["field"] == "keywords" for r in records))


# ==========================================================================
class TestAdsMetrics(unittest.TestCase):
    """指标口径：整个系统共用这一份计算，错了会全盘错。"""

    def test_metrics_math(self) -> None:
        m = ads_domain._metrics(10000, 40, 32.0, 4, 5, 160.0)
        self.assertAlmostEqual(m["ctr"], 0.004, places=6)
        self.assertAlmostEqual(m["cvr"], 0.1, places=6)
        self.assertAlmostEqual(m["cpc"], 0.8, places=6)
        self.assertAlmostEqual(m["acos"], 0.2, places=6)
        self.assertAlmostEqual(m["roas"], 5.0, places=6)

    def test_metrics_zero_sales_gives_none_not_infinity(self) -> None:
        m = ads_domain._metrics(5000, 20, 15.0, 0, 0, 0.0)
        self.assertIsNone(m["acos"], "零销售时 ACOS 必须为 None，不能是 inf（会污染排序与 JSON）")
        self.assertIsNone(m["cpa"])
        self.assertEqual(m["cvr"], 0.0)

    def test_week_bounds_is_monday_to_sunday(self) -> None:
        start, end = ads_domain.week_bounds("2026-10-07")   # 周三
        self.assertEqual(start, "2026-10-05")
        self.assertEqual(end, "2026-10-11")

    def test_default_report_week_is_always_complete(self) -> None:
        start = ads_domain.default_report_week()
        end = ads_domain.week_bounds(start)[1]
        from datetime import datetime
        self.assertLess(datetime.fromisoformat(end).date(), date.today(),
                        "默认报告周必须已经结束，否则是拿三天数据跟整周比")

    # 回归：`p_end, _ = week_bounds(p_start)` 把"周一"当成了"对比周的周日"，
    # 于是环比变成"整周 vs 单日"，报告里出现了 +617% 这种荒谬数字。
    def test_wow_compares_full_weeks(self) -> None:
        self.assertEqual(ads_domain.prev_week("2026-10-05"), "2026-09-28")
        p_start = ads_domain.prev_week("2026-10-05")
        p_end = ads_domain.week_bounds(p_start)[1]
        self.assertEqual(p_end, "2026-10-04", "对比周必须是完整 7 天（周一~周日）")
        span = (date.fromisoformat(p_end) - date.fromisoformat(p_start)).days
        self.assertEqual(span, 6, "对比周跨度应为 6 天（含首尾 7 天）")


# ==========================================================================
class TestAdsDiagnosis(TempWorkspace):
    """诊断规则引擎：每条结论都必须有依据、不猜。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=30, days_ads=21)
        cls.week = ads_domain.default_report_week()
        cls.metrics = ads_domain.weekly_metrics(cls.week)
        cls.findings = ads_domain.diagnose(cls.metrics)

    def test_metrics_have_data(self) -> None:
        self.assertGreater(self.metrics["total"]["spend"], 0)
        self.assertGreater(len(self.metrics["by_campaign"]), 0)

    def test_findings_are_well_formed(self) -> None:
        self.assertTrue(self.findings)
        for f in self.findings:
            with self.subTest(rule=f["rule"]):
                self.assertIn(f["severity"], ("critical", "warning", "opportunity", "info"))
                self.assertIn(f["feeds"], ("listing", "ads", "inventory"))
                self.assertTrue(f["diagnosis"] and f["action"])
                self.assertIsInstance(f["evidence"], dict)
                self.assertTrue(f["evidence"], "每条结论都必须挂可核对的证据数字")

    def test_findings_sorted_by_severity(self) -> None:
        order = {"critical": 0, "warning": 1, "opportunity": 2, "info": 3}
        keys = [order[f["severity"]] for f in self.findings]
        self.assertEqual(keys, sorted(keys))

    def test_zero_order_needs_enough_clicks(self) -> None:
        """回归：点击量不足时不得下"只花钱不出单"的结论。"""
        for f in self.findings:
            if f["rule"] == "zero_order_burn":
                self.assertGreaterEqual(f["evidence"]["clicks"],
                                        config.Biz.MIN_CLICKS_FOR_JUDGE)

    def test_low_aov_rule_only_when_cpc_at_floor(self) -> None:
        """回归：客单价问题只应在 CPC 已到底价时判定，否则会把"买贵了"误判成"卖便宜了"。"""
        for f in self.findings:
            if f["rule"] == "low_aov_unprofitable":
                self.assertLessEqual(f["evidence"]["cpc"], config.Biz.CPC_FLOOR_ZONE)
                self.assertLess(f["evidence"]["aov"], f["evidence"]["required_aov"])

    def test_budget_capped_opportunity_has_evidence(self) -> None:
        for f in self.findings:
            if f["rule"] == "budget_capped_scale":
                self.assertGreaterEqual(f["evidence"]["budget_utilization"],
                                        config.Biz.BUDGET_CAP_RATIO)

    def test_listing_targets_are_real_skus(self) -> None:
        from saviland.pipelines import weekly as weekly_pipe
        targets = weekly_pipe.listing_targets(self.findings)
        for sku in targets:
            self.assertIsNotNone(catalog_domain.get_product(sku),
                                 f"闭环指向的 SKU 必须真实存在：{sku}")

    # 回归：账户 ACOS 29.2% 时曾报"健康"，因为判定只比了预警线(35%)、漏了目标线(28%)
    def test_account_rule_respects_target_line(self) -> None:
        acos = self.metrics["total"]["acos"]
        rules = {f["rule"] for f in self.findings}
        if acos is not None and acos > config.Biz.TARGET_ACOS:
            self.assertNotIn("account_acos_healthy", rules,
                             "ACOS 高于目标时不应报健康")

    def test_diagnostics_summary(self) -> None:
        s = ads_domain.diagnostics_summary(self.findings)
        self.assertEqual(s["total"], len(self.findings))
        self.assertGreaterEqual(s["waste_usd"], 0)
        self.assertEqual(s["listing_actions"] + s["ads_actions"],
                         sum(1 for f in self.findings if f["feeds"] in ("listing", "ads")))


# ==========================================================================
class TestBusinessDomain(TempWorkspace):
    """电商后台业务层：库存 / 订单 / 商品。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        if not (db.scalar("SELECT COUNT(*) FROM products") or 0):
            seed.seed_all(force=True, days_orders=30, days_ads=14)

    def test_products_seeded_from_catalog(self) -> None:
        products = catalog_domain.list_products()
        self.assertEqual(len(products), 20)
        self.assertTrue(all(p["price_usd"] > 0 for p in products))
        self.assertTrue(all(p["cost_usd"] > 0 for p in products))

    def test_margin_between_zero_and_one(self) -> None:
        for p in catalog_domain.list_products():
            with self.subTest(sku=p["sku"]):
                self.assertTrue(0 < catalog_domain.margin_of(p) < 1)

    def test_product_attributes_roundtrip(self) -> None:
        product = catalog_domain.get_product("SV-UVL-48W")
        self.assertIsNotNone(product)
        self.assertTrue(product["selling_points"], "卖点应能从 product_attributes 读回")
        self.assertTrue(product["specs"], "规格应能读回")

    def test_listing_copy_versioning(self) -> None:
        sku = "SV-UVL-48W"
        before = catalog_domain.next_copy_version(sku)
        catalog_domain.save_listing_copy({"sku": sku, "title": "t1", "bullets": ["a"],
                                          "keywords": ["k"], "engine": "test"})
        catalog_domain.save_listing_copy({"sku": sku, "title": "t2", "bullets": ["b"],
                                          "keywords": ["k2"], "engine": "test"})
        after = catalog_domain.next_copy_version(sku)
        self.assertEqual(after, before + 2)
        latest = catalog_domain.latest_listing_copy(sku)
        self.assertEqual(latest["title"], "t2")
        self.assertEqual(latest["version"], after - 1)

    def test_revenue_excludes_refunds(self) -> None:
        o = orders_domain.sales_summary(30)
        self.assertGreater(o["revenue"], 0)
        self.assertGreater(o["units"], 0)
        self.assertGreaterEqual(o["refund_units"], 0)
        # 毛利 = 营收 - 成本，口径必须自洽
        self.assertAlmostEqual(o["gross_profit"], o["revenue"] - o["cogs"], places=2)

    def test_velocity_positive_for_hot_skus(self) -> None:
        self.assertGreater(inv_domain.sales_velocity("SV-UVL-48W", 14), 0)

    def test_available_days_math(self) -> None:
        sku = "SV-UVL-48W"
        velocity = inv_domain.sales_velocity(sku, 14)
        available = inv_domain.stock_of(sku)["available"]
        expected = round(available / velocity, 1)
        self.assertAlmostEqual(inv_domain.available_days(sku, 14), expected, places=1)

    def test_restock_advice_levels(self) -> None:
        advice = inv_domain.restock_advice()
        self.assertEqual(len(advice), 20)
        valid = {"stockout_risk", "overstock", "healthy", "no_sales"}
        for a in advice:
            with self.subTest(sku=a["sku"]):
                self.assertIn(a["level"], valid)
                self.assertGreaterEqual(a["suggest_qty"], 0)
                if a["level"] == "stockout_risk":
                    self.assertLessEqual(a["available_days"], config.Biz.STOCKOUT_WARN_DAYS)
                if a["level"] == "overstock":
                    self.assertGreaterEqual(a["available_days"], config.Biz.OVERSTOCK_DAYS)


# ==========================================================================
class TestSeedDeterminism(TempWorkspace):
    """仿真数据必须可复现：同一台机器跑两次，周报要一模一样。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=14, days_ads=14)
        cls.first = ads_domain.weekly_metrics(ads_domain.default_report_week())["total"]
        cls.first_orders = db.scalar("SELECT COUNT(*) FROM orders")
        seed.seed_all(force=True, days_orders=14, days_ads=14)
        cls.second = ads_domain.weekly_metrics(ads_domain.default_report_week())["total"]

    def test_same_seed_same_totals(self) -> None:
        for key in ("spend", "sales", "orders", "clicks", "impressions"):
            with self.subTest(key=key):
                self.assertEqual(self.first[key], self.second[key],
                                 "固定随机种子的数据必须完全可复现")

    def test_order_count_stable(self) -> None:
        self.assertEqual(self.first_orders, db.scalar("SELECT COUNT(*) FROM orders"))

    def test_no_duplicate_campaign_rows(self) -> None:
        rows = db.query("SELECT campaign_id, COUNT(*) AS n FROM ad_campaigns GROUP BY campaign_id")
        self.assertTrue(all(r["n"] == 1 for r in rows))


# ==========================================================================
class TestLLMHelpers(unittest.TestCase):
    """LLM 输出的脏数据解析（不联网）。"""

    def test_extract_plain_json(self) -> None:
        self.assertEqual(llm.extract_json('{"a": 1}'), {"a": 1})

    def test_extract_from_code_fence(self) -> None:
        text = '好的，这是结果：\n```json\n{"title": "x", "bullets": ["a"]}\n```\n希望有帮助。'
        self.assertEqual(llm.extract_json(text), {"title": "x", "bullets": ["a"]})

    def test_extract_from_noisy_text(self) -> None:
        self.assertEqual(llm.extract_json('blah {"k": [1,2]} trailing'), {"k": [1, 2]})

    def test_extract_array(self) -> None:
        self.assertEqual(llm.extract_json('[{"a": 1}]'), [{"a": 1}])

    def test_extract_returns_none_on_garbage(self) -> None:
        self.assertIsNone(llm.extract_json("完全不是 JSON"))
        self.assertIsNone(llm.extract_json(""))

    def test_offline_engine_never_calls_network(self) -> None:
        res = llm.chat("你好", engine="offline")
        self.assertFalse(res.ok)
        self.assertEqual(res.engine, "offline")
        self.assertIn("离线", res.note)


# ==========================================================================
class TestHeroRendering(TempWorkspace):
    """主图渲染：遮罩模式、合规自检、缺图兜底。"""

    PRODUCT = {
        "sku": "TEST-001", "title_cn": "测试商品", "title_en": "Test Product",
        "category": "美甲灯", "price_usd": 19.99, "cost_usd": 6.0,
        "specs": {"功率": "48W", "灯珠": "36 颗"}, "selling_points": ["快", "冷光", "耐用"],
        "hero_image": "",
    }
    COPY = {
        "hero_headline": "Fast Curing in 30 Seconds",
        "hero_subhead": "36 LED Beads, Low Heat Mode",
        "hero_badge": "48W · 36 LEDs",
        "cta": "Salon Results at Home",
        "bullets": ["Quick curing", "Even coverage", "Low heat"],
    }

    def test_render_with_missing_image_uses_placeholder(self) -> None:
        img = hero_render.render_variant(self.PRODUCT, hero_render.Copy.from_dict(self.COPY, self.PRODUCT), 0)
        self.assertEqual(img.size, (config.Biz.HERO_SIZE, config.Biz.HERO_SIZE))
        self.assertEqual(img.mode, "RGB")

    # 回归：`_drop_shadow` 把 RGB 商品图当遮罩传给 PIL，抛 "bad transparency mask"。
    # 症状是"主图时好时坏"—— 商品图缺失时（走 RGBA 占位图）能跑，图一到位就崩。
    def test_render_with_rgb_product_image(self) -> None:
        path = config.ASSETS_DIR / "TEST-001.png"
        Image.new("RGB", (1600, 1600), (240, 240, 245)).save(path)   # RGB，无 alpha
        product = dict(self.PRODUCT, hero_image=str(path))
        for variant in (0, 1, 2):
            with self.subTest(variant=variant):
                img = hero_render.render_variant(
                    product, hero_render.Copy.from_dict(self.COPY, product), variant)
                self.assertEqual(img.size, (config.Biz.HERO_SIZE, config.Biz.HERO_SIZE))

    def test_render_with_rgba_product_image(self) -> None:
        path = config.ASSETS_DIR / "TEST-002.png"
        Image.new("RGBA", (1200, 900), (255, 255, 255, 0)).save(path)
        product = dict(self.PRODUCT, sku="TEST-002", hero_image=str(path))
        img = hero_render.render_variant(product, hero_render.Copy.from_dict(self.COPY, product), 0)
        self.assertEqual(img.mode, "RGB")

    def test_render_with_corrupt_image_does_not_crash(self) -> None:
        path = config.ASSETS_DIR / "TEST-003.png"
        path.write_bytes(b"not a real png")
        product = dict(self.PRODUCT, sku="TEST-003", hero_image=str(path))
        img = hero_render.render_variant(product, hero_render.Copy.from_dict(self.COPY, product), 0)
        self.assertEqual(img.size, (config.Biz.HERO_SIZE, config.Biz.HERO_SIZE))

    def test_compose_and_compliance_check(self) -> None:
        product = dict(self.PRODUCT, sku="TEST-004", hero_image="")
        out = hero_render.compose_hero(product, self.COPY, variant=0)
        self.assertTrue(out.is_file())
        report = hero_render.compliance_check(out)
        self.assertTrue(report["square"])
        self.assertGreaterEqual(min(Image.open(out).size), 1000)
        self.assertTrue(report["white_corners"], "clean 版式四角必须是纯白（亚马逊主图要求）")
        self.assertTrue(report["passed"], f"clean 版式应通过合规自检：{report}")

    def test_long_text_does_not_overflow(self) -> None:
        """超长文案必须被截断装饰，而不是画到画布外面。"""
        copy = dict(self.COPY,
                    hero_headline="A Very Long Headline " * 6,
                    hero_subhead="Also a very long subheadline " * 5,
                    hero_badge="B" * 80)
        product = dict(self.PRODUCT, sku="TEST-005", hero_image="")
        for variant in (0, 1, 2):
            with self.subTest(variant=variant):
                img = hero_render.render_variant(product, hero_render.Copy.from_dict(copy, product), variant)
                self.assertEqual(img.size, (config.Biz.HERO_SIZE, config.Biz.HERO_SIZE))

    def test_contact_sheet(self) -> None:
        paths = [hero_render.compose_hero(dict(self.PRODUCT, sku=f"TEST-{i:03d}", hero_image=""),
                                          self.COPY, variant=0) for i in range(3)]
        sheet = hero_render.contact_sheet(paths, config.OUTPUTS_DIR / "sheet.png", cols=3, cell=300)
        self.assertTrue(sheet.is_file())
        self.assertGreater(sheet.stat().st_size, 1000)


# ==========================================================================
class TestRagBridgePure(unittest.TestCase):
    """RAG 桥接层里**不需要联网**的部分。"""

    def test_term_label_handles_tuples(self) -> None:
        """回归：底座 build_terms 返回 (词, 权重)，直接 join 会抛
        `TypeError: sequence item 0: expected str instance, tuple found`。"""
        from saviland import rag_bridge
        self.assertEqual(rag_bridge._term_label(("美甲灯", 0.6)), "美甲灯")
        self.assertEqual(rag_bridge._term_label(["灯", 1.0]), "灯")
        self.assertEqual(rag_bridge._term_label("词"), "词")
        self.assertEqual(rag_bridge._term_label(123), "123")

    def test_merge_kb_line_ranges_are_correct(self) -> None:
        """回归：区间索引曾按 `len(lines)`（列表元素数）计数，而 markdown 正文是
        一个含大量换行的字符串，于是 1111 行的知识库被算成"每文件 6 行"，
        检索结果的来源全部退化成"未定位到具体文件"。"""
        from saviland import rag_bridge
        merged, count = rag_bridge.merge_kb(force=True)
        lines = Path(merged).read_text(encoding="utf-8").splitlines()
        ranges = rag_bridge._load_ranges()
        self.assertGreater(count, 0)
        self.assertEqual(len(ranges), count)
        for r in ranges:
            with self.subTest(file=r["file"]):
                self.assertLess(r["start"], r["end"] - 5,
                                "每个知识库文件的区间必须远超 6 行（否则就是计数口径错了）")
        # 区间必须首尾衔接、覆盖整个文件
        self.assertEqual(ranges[0]["start"], 1)
        for a, b in zip(ranges, ranges[1:]):
            self.assertEqual(a["end"] + 1, b["start"], "区间之间不能有缝也不能重叠")
        self.assertLessEqual(ranges[-1]["end"], len(lines) + 1)

    def test_source_of_uses_line_range(self) -> None:
        from saviland import rag_bridge
        rag_bridge.merge_kb(force=True)
        rag_bridge._RANGES = rag_bridge._load_ranges()
        target = rag_bridge._RANGES[2]
        mid = (target["start"] + target["end"]) // 2
        self.assertEqual(rag_bridge._source_of("正文里没有文件标记", mid), target["file"])

    def test_source_of_falls_back_to_marker(self) -> None:
        from saviland import rag_bridge
        rag_bridge._RANGES = []
        text = f"{rag_bridge.SECTION_MARK}\n# 【知识库文件】05-合规与禁用词.md\n内容"
        self.assertEqual(rag_bridge._source_of(text, -1), "05-合规与禁用词.md")

    def test_rewrite_prompt_is_domain_adapted(self) -> None:
        """底座的改写提示词是 IT 服务台专用的，桥接层必须换成电商域。"""
        from saviland import rag_bridge
        self.assertIn("美甲", rag_bridge.ECOM_REWRITE_PROMPT)
        self.assertNotIn("IT 服务台", rag_bridge.ECOM_REWRITE_PROMPT)
        # JSON 契约必须与底座一致，否则底座的容错解析会失效
        self.assertIn('"query"', rag_bridge.ECOM_REWRITE_PROMPT)
        self.assertIn('"keywords"', rag_bridge.ECOM_REWRITE_PROMPT)

    def test_product_sections_include_sku(self) -> None:
        from saviland import rag_bridge
        text = rag_bridge._product_sections()
        self.assertIn("SV-UVL-48W", text)
        self.assertIn("【知识库文件】08-产品档案.md", text)


# ==========================================================================
class TestListingOffline(TempWorkspace):
    """链路 1 的离线兜底必须真的能出可用产物（无 Key 环境下也交得出东西）。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        if not (db.scalar("SELECT COUNT(*) FROM products") or 0):
            seed.seed_all(force=True, days_orders=14, days_ads=14)

    def test_offline_copy_is_complete(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        product = catalog_domain.get_product("SV-UVL-48W")
        copy = listing_pipe.offline_copy(product)
        for key in ("title", "hero_headline", "hero_subhead", "hero_badge",
                    "cta", "bullets", "description", "keywords"):
            with self.subTest(key=key):
                self.assertTrue(copy[key], f"离线兜底也必须给出 {key}")
        self.assertEqual(len(copy["bullets"]), 5)
        self.assertLessEqual(len(copy["keywords"]), 15)
        self.assertGreaterEqual(len(copy["hero_copy_variants"]), 2)

    def test_normalize_fills_missing_fields(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        product = catalog_domain.get_product("SV-GEL-CAT")
        out = listing_pipe._normalize({"title": "Only a title"}, product)
        self.assertEqual(len(out["bullets"]), 5)
        self.assertTrue(out["hero_headline"])
        self.assertTrue(out["keywords"])

    def test_generate_copy_offline_end_to_end(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        res = listing_pipe.generate_one("SV-TOOL-DRILL", variants=1, engine="offline",
                                        use_rag=False)
        self.assertTrue(res.ok)
        self.assertTrue(res.degraded, "offline 引擎必须被标记为降级，不能假装成功")
        self.assertIsNotNone(res.version)
        self.assertEqual(len(res.hero_images), 1)
        self.assertTrue(Path(res.hero_images[0]).is_file())
        # 运行档案必须留痕
        from saviland import runs
        self.assertIn(res.run_id, [r["run_id"] for r in runs.recent_runs("listing", 20)])

    def test_variant_copy_falls_back(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        base = {"hero_headline": "H", "hero_subhead": "S", "hero_badge": "B",
                "hero_copy_variants": []}
        merged = listing_pipe.copy_for_variant(base, 2)
        self.assertEqual(merged["hero_headline"], "H", "没有变体时应回落到主方案而不是报错")

    # 回归：文档/作品集引用的样例曾用 `batch_latest.csv`，而批量通常只覆盖几个 SKU，
    # 于是"20 个 SKU 都跑通了"在样例里变成 4 行。同类问题还有主图总览图被覆盖成 4 个。
    # 根因都是把"最近一次"当成了"全部现状"。
    def test_export_all_latest_covers_every_sku_with_copy(self) -> None:
        import csv as _csv

        from saviland.pipelines import listing as listing_pipe

        for sku in ("SV-UVL-48W", "SV-GEL-CAT", "SV-TOOL-DRILL"):
            listing_pipe.generate_one(sku, variants=1, engine="offline", use_rag=False)
        path = listing_pipe.export_all_latest()

        rows = list(_csv.reader(path.open(encoding="utf-8-sig")))
        header, body = rows[0], rows[1:]
        expected = [r["sku"] for r in db.query(
            "SELECT DISTINCT sku FROM listing_copies ORDER BY sku")]
        self.assertEqual(len(body), len(expected),
                         "导出必须覆盖**全部**有文案的 SKU，而不是最近一次批量")
        self.assertEqual([r[0] for r in body], expected)
        self.assertEqual(header[0], "sku")
        self.assertEqual(len(header), len(body[0]), "每行列数必须与表头一致")
        for row in body:
            self.assertTrue(row[5], f"{row[0]} 的标题不应为空")

    def test_export_all_latest_picks_newest_version(self) -> None:
        import csv as _csv

        from saviland.pipelines import listing as listing_pipe
        listing_pipe.generate_one("SV-GEL-TOPC", variants=1, engine="offline", use_rag=False)
        listing_pipe.generate_one("SV-GEL-TOPC", variants=1, engine="offline", use_rag=False)
        latest = catalog_domain.latest_listing_copy("SV-GEL-TOPC")
        rows = {r[0]: r for r in _csv.reader(
            listing_pipe.export_all_latest().open(encoding="utf-8-sig"))}
        self.assertEqual(int(rows["SV-GEL-TOPC"][1]), latest["version"],
                         "导出的必须是该 SKU 的最新版本")


# ==========================================================================
class TestWeeklyOffline(TempWorkspace):
    """链路 2 的离线兜底：报告必须完整产出（MD + HTML + Excel + 图表）。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=21, days_ads=21)
        from saviland.pipelines import weekly as weekly_pipe
        cls.res = weekly_pipe.generate_weekly(engine="offline", use_rag=False)

    def test_report_generated(self) -> None:
        self.assertTrue(self.res.ok, f"离线周报应生成成功：{self.res.note}")
        self.assertGreater(len(self.res.findings), 0)
        self.assertTrue(self.res.degraded)

    def test_all_artifacts_exist(self) -> None:
        for name, path in (("markdown", self.res.report_md), ("html", self.res.report_html),
                           ("excel", self.res.report_xlsx)):
            with self.subTest(artifact=name):
                p = Path(path)
                self.assertTrue(p.is_file(), f"{name} 未生成：{p}")
                self.assertGreater(p.stat().st_size, 800)

    def test_markdown_has_all_sections(self) -> None:
        md = Path(self.res.report_md).read_text(encoding="utf-8")
        for section in ("执行摘要", "关键指标总览", "问题诊断与处置建议", "活动明细排名",
                        "环比变化与流量质量分诊", "库存与广告的联动风险",
                        "下周行动清单", "链路闭环", "附录：数据口径"):
            with self.subTest(section=section):
                self.assertIn(section, md)

    def test_html_is_self_contained(self) -> None:
        html = Path(self.res.report_html).read_text(encoding="utf-8")
        self.assertIn("<!DOCTYPE html>", html)
        # 图片内嵌成 base64，才能当一个文件直接转发
        self.assertIn("data:image/png;base64,", html)

    def test_charts_generated(self) -> None:
        self.assertGreaterEqual(len(self.res.charts), 4)
        for p in self.res.charts:
            with self.subTest(chart=Path(p).stem):
                self.assertTrue(Path(p).is_file())
                self.assertGreater(Path(p).stat().st_size, 3000)

    def test_excel_has_expected_sheets(self) -> None:
        from openpyxl import load_workbook
        wb = load_workbook(self.res.report_xlsx)
        self.assertEqual(
            wb.sheetnames,
            ["账户总览", "活动明细", "SKU明细", "每日趋势", "诊断清单", "行动清单"],
        )

    def test_ai_sections_present(self) -> None:
        ai = self.res.ai
        self.assertTrue(ai.get("executive_summary"))
        self.assertTrue(ai.get("action_plan"))
        for a in ai["action_plan"]:
            self.assertIn("action", a)

    def test_run_recorded(self) -> None:
        from saviland import runs
        run = next((r for r in runs.recent_runs("weekly", 10) if r["run_id"] == self.res.run_id), None)
        self.assertIsNotNone(run, "周报运行必须写入 generation_runs")
        self.assertEqual(run["status"], "degraded")
        self.assertTrue(runs.artifacts_of(self.res.run_id))


# ==========================================================================
class TestAttributionMapping(unittest.TestCase):
    """把模型返回的归因映射回 findings。这块张冠李戴过一次，必须钉死。"""

    FINDINGS = [
        {"rule": "high_acos_low_cvr", "diagnosis": "打磨机 ACOS 142% 且转化率仅 1.4%",
         "subject_name": "打磨机"},
        {"rule": "high_acos_low_cvr", "diagnosis": "专业套装 ACOS 50% 且转化率 6.2%",
         "subject_name": "专业套装"},
        {"rule": "zero_order_burn", "diagnosis": "水钻饰品盒 227 次点击零订单",
         "subject_name": "水钻饰品盒"},
    ]

    def test_shared_rule_does_not_leak_attribution(self) -> None:
        """回归：最初按 `rule` 做键映射，而两条 finding 共用同一条规则，
        字典合并后只剩最后一条 —— "打磨机"的归因栏里写的是"专业套装"的原因。
        张冠李戴的归因比没有归因更糟，因为老板会照着错的结论去改。"""
        from saviland.pipelines import weekly as weekly_pipe
        ai = {"attributions": [
            {"id": 1, "insight": "广泛匹配引入了 283 次找电动指甲锉的点击，落地页不匹配"},
            {"id": 2, "insight": "品牌广告竞价被抬高到 2.09 美元，超出这个客单价的承受范围"},
        ]}
        attr = weekly_pipe.attributions_of(ai, self.FINDINGS)
        self.assertIn(1, attr)
        self.assertIn(2, attr)
        self.assertIn("电动指甲锉", attr[1])
        self.assertIn("品牌广告", attr[2])
        self.assertNotIn("品牌广告", attr[1], "第 1 条的归因绝不能串成第 2 条的内容")

    def test_restating_the_symptom_is_filtered(self) -> None:
        from saviland.pipelines import weekly as weekly_pipe
        ai = {"attributions": [
            {"id": 3, "insight": "水钻饰品盒 227 次点击零订单"},
            # 好归因：不复述现象，点出机制（泛意图流量）
            {"id": 1, "insight": "广泛匹配把泛词也放进来了，多数点击来自找电动指甲锉的买家，落地页不匹配"},
        ]}
        attr = weekly_pipe.attributions_of(ai, self.FINDINGS)
        self.assertNotIn(3, attr, "只是复述现象的归因没有信息量，应被过滤掉")
        self.assertIn(1, attr)

    def test_rule_key_fallback_still_works(self) -> None:
        """离线兜底给的是 rule 键（没有 id），必须兼容。"""
        from saviland.pipelines import weekly as weekly_pipe
        ai = {"attributions": [{"rule": "zero_order_burn",
                                "insight": "自动投放把大词全打开了，227 次点击来自泛意图流量"}]}
        attr = weekly_pipe.attributions_of(ai, self.FINDINGS)
        self.assertIn(3, attr)


# ==========================================================================
class TestAdsCsvImport(TempWorkspace):
    """导入亚马逊广告报表 —— 「可切换真实数据」这条设计线的落点。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=14, days_ads=14)
        cls.csv = cls.tmp / "amazon_report.csv"

    def _write(self, text: str) -> Path:
        self.csv.write_text(text, encoding="utf-8-sig")
        return self.csv

    def test_import_with_english_headers(self) -> None:
        """亚马逊后台导出的英文表头（含空格与 7 Day Total 前缀）。"""
        path = self._write(
            "Date,Campaign ID,Campaign Name,Advertised SKU,Impressions,Clicks,Spend,"
            "7 Day Total Orders,7 Day Total Units,7 Day Total Sales\n"
            "2019-01-07,C001-SV-UVL-48W,Saviland 48W,SV-UVL-48W,12000,60,44.20,6,7,180.50\n"
            "2019-01-08,C001-SV-UVL-48W,Saviland 48W,SV-UVL-48W,11000,55,40.10,5,6,155.00\n"
        )
        out = seed.import_ads_csv(path)
        self.assertEqual(out["rows"], 2)
        self.assertEqual(out["skipped"], 0)
        self.assertIn("date", out["mapped_columns"])
        row = db.query_one("SELECT * FROM ad_daily WHERE date = '2019-01-07'")
        self.assertIsNotNone(row, "导入的行必须真的落库")
        self.assertEqual(row["sku"], "SV-UVL-48W")
        self.assertAlmostEqual(row["sales_usd"], 180.50, places=2)

    def test_import_with_chinese_headers(self) -> None:
        """中文表头也要能吃掉（列名做了中英文模糊映射）。"""
        path = self._write(
            "日期,广告活动 ID,广告活动名称,商品 SKU,曝光量,点击量,花费,订单量,销售量,销售额\n"
            "2019-02-01,C001-SV-UVL-48W,Saviland 48W,SV-UVL-48W,9000,45,30.00,3,4,120.00\n"
        )
        out = seed.import_ads_csv(path)
        self.assertEqual(out["rows"], 1)

    def test_money_symbols_and_commas_are_parsed(self) -> None:
        path = self._write(
            "Date,Campaign ID,SKU,Impressions,Clicks,Spend,Orders,Units,Sales\n"
            '2019-03-01,C001-SV-UVL-48W,SV-UVL-48W,"10,500",50,"$38.75",4,5,"$1,240.00"\n'
        )
        seed.import_ads_csv(path)
        row = db.query_one("SELECT * FROM ad_daily WHERE date = '2019-03-01'")
        self.assertEqual(row["impressions"], 10500)
        self.assertAlmostEqual(row["spend_usd"], 38.75, places=2)
        self.assertAlmostEqual(row["sales_usd"], 1240.00, places=2)

    def test_slash_and_compact_dates(self) -> None:
        path = self._write(
            "Date,Campaign ID,SKU,Impressions,Clicks,Spend,Orders,Units,Sales\n"
            "2019/04/05,C001-SV-UVL-48W,SV-UVL-48W,100,5,1.00,0,0,0.00\n"
            "20190406,C001-SV-UVL-48W,SV-UVL-48W,100,5,1.00,0,0,0.00\n"
        )
        seed.import_ads_csv(path)
        self.assertIsNotNone(db.query_one("SELECT 1 FROM ad_daily WHERE date = '2019-04-05'"))
        self.assertIsNotNone(db.query_one("SELECT 1 FROM ad_daily WHERE date = '2019-04-06'"))

    def test_dirty_rows_are_skipped_not_fatal(self) -> None:
        """单行脏数据不该毁掉整份报表 —— 真实导出里空行/缺列很常见。"""
        path = self._write(
            "Date,Campaign ID,SKU,Impressions,Clicks,Spend,Orders,Units,Sales\n"
            "2019-05-01,C001-SV-UVL-48W,SV-UVL-48W,100,5,1.00,0,0,0.00\n"
            ",,,0,0,0,0,0,0\n"
            "2019-05-02,,,100,5,1.00,0,0,0.00\n"
            "2019-05-03,C001-SV-UVL-48W,SV-UVL-48W,100,5,1.00,0,0,0.00\n"
        )
        out = seed.import_ads_csv(path)
        self.assertEqual(out["rows"], 2, "两行有效数据必须都进来")
        self.assertEqual(out["skipped"], 2, "两行脏数据应被跳过而不是报错")

    def test_missing_required_columns_raises_clear_error(self) -> None:
        path = self._write("Foo,Bar\n1,2\n")
        with self.assertRaises(ValueError) as ctx:
            seed.import_ads_csv(path)
        self.assertIn("date", str(ctx.exception))

    def test_reimport_is_idempotent(self) -> None:
        """(date, campaign_id) 有唯一约束，重复导入不该翻倍。"""
        path = self._write(
            "Date,Campaign ID,SKU,Impressions,Clicks,Spend,Orders,Units,Sales\n"
            "2019-06-01,C001-SV-UVL-48W,SV-UVL-48W,100,5,2.00,1,1,20.00\n"
        )
        seed.import_ads_csv(path)
        seed.import_ads_csv(path)
        n = db.scalar("SELECT COUNT(*) FROM ad_daily WHERE date = '2019-06-01'")
        self.assertEqual(n, 1)
        row = db.query_one("SELECT spend_usd FROM ad_daily WHERE date = '2019-06-01'")
        self.assertAlmostEqual(row["spend_usd"], 2.00, places=2)


class TestWhatIf(TempWorkspace):
    """"现场改数据验证规则"必须是**可执行的命令**，不是说明。

    ★ 回归：演示说明里写了"我可以现场改数据 —— 把花费调高、订单清零，
    '只花钱不出单'的结论立刻出现"，但很久以来**没有对应命令**。
    这个类钉住三件事：改得动、结论真的会变、改完能还原。
    """

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=21, days_ads=21)

    def test_resolve_campaign_by_id_prefix_and_name(self) -> None:
        from saviland import whatif
        camp = whatif._resolve_campaign("C002")
        self.assertTrue(camp["campaign_id"].startswith("C002"))
        # 写完整名称也能定位（名称里带 SKU 和投放方式，是唯一的）
        self.assertEqual(whatif._resolve_campaign(camp["name"])["campaign_id"],
                         camp["campaign_id"])

    def test_ambiguous_fragment_lists_candidates(self) -> None:
        """模糊到多个活动时必须报错并给出候选，不能随便挑一个 ——
        现场演示时"改错了活动"比"报错"难看得多。"""
        from saviland import whatif
        with self.assertRaises(ValueError) as ctx:
            whatif._resolve_campaign("手动精准")
        msg = str(ctx.exception)
        self.assertIn("多个活动", msg)
        self.assertIn("C0", msg, "报错里要带上候选活动 id")
        # 候选里不该出现"和本次无关"的活动名之外的东西（至少长度是合理的）
        self.assertGreater(len(msg), 20)

    def test_unknown_campaign_raises_with_hint(self) -> None:
        from saviland import whatif
        with self.assertRaises(ValueError) as ctx:
            whatif._resolve_campaign("不存在的活动")
        self.assertIn("没找到", str(ctx.exception))

    def _pick_target(self) -> tuple[str, str] | None:
        """挑一个"当前有单、且没有 zero_order_burn"的活动，最适合演示。"""
        from saviland.domain import ads as ads_domain
        week = ads_domain.default_report_week()
        m = ads_domain.weekly_metrics(week)
        rules: dict[str, set[str]] = {}
        for f in ads_domain.diagnose(m):
            rules.setdefault(f.get("subject"), set()).add(f["rule"])
        for c in m["by_campaign"]:
            if c["orders"] > 0 and "zero_order_burn" not in rules.get(c["campaign_id"], set()):
                return c["campaign_id"], week
        return None

    def test_zeroing_orders_triggers_zero_order_burn(self) -> None:
        from saviland import whatif
        picked = self._pick_target()
        if not picked:
            self.skipTest("测试数据里没有「有单且健康」的活动可改")
        campaign, week = picked

        res = whatif.whatif(campaign=campaign, orders=0, week=week)
        self.assertEqual(res["after"]["orders"], 0)
        self.assertEqual(res["after"]["sales"], 0.0, "0 单却有销售额是不自洽的数据")
        rules = {f["rule"] for f in res["new_findings"]}
        self.assertIn("zero_order_burn", rules,
                      f"把 {campaign} 的订单清零后应触发 zero_order_burn，实际新增：{rules}")
        self.assertEqual(res["after"]["clicks"], res["before"]["clicks"],
                         "只改订单时点击量不该变（除非为了跨过判断门槛小幅补齐）")

    def test_restores_data_by_default(self) -> None:
        """默认必须还原 —— 演示完不能把库改坏。"""
        from saviland import whatif
        picked = self._pick_target()
        if not picked:
            self.skipTest("没有可改的活动")
        campaign, week = picked

        before_orders = whatif.whatif(campaign=campaign, orders=0, week=week)["before"]["orders"]
        row = db.query_one(
            "SELECT COALESCE(SUM(orders),0) AS n FROM ad_daily a "
            "WHERE a.campaign_id = ? AND a.date BETWEEN ? AND ?",
            (campaign, *ads_week_bounds(week)),
        )
        self.assertEqual(int(row["n"]), before_orders, "默认应还原到改动前的订单数")

    def test_keep_leaves_change_in_place(self) -> None:
        from saviland import whatif
        picked = self._pick_target()
        if not picked:
            self.skipTest("没有可改的活动")
        campaign, week = picked
        try:
            whatif.whatif(campaign=campaign, orders=0, week=week, keep=True)
            row = db.query_one(
                "SELECT COALESCE(SUM(orders),0) AS n FROM ad_daily a "
                "WHERE a.campaign_id = ? AND a.date BETWEEN ? AND ?",
                (campaign, *ads_week_bounds(week)),
            )
            self.assertEqual(int(row["n"]), 0, "加了 keep 就该保留改动")
        finally:
            seed.reseed_ads()


def ads_week_bounds(week: str) -> tuple[str, str]:
    from saviland.domain import ads as ads_domain
    return ads_domain.week_bounds(week)


# ==========================================================================
class TestInterpreterGuard(unittest.TestCase):
    """解释器守卫：用错 python 时要自动换，而不是甩 ModuleNotFoundError。

    ★ 回归：机器上默认 `python` 是 3.14、而依赖装在 3.12 的底座 venv 里，
    直接 `python tests/run_all.py` 会得到 4 个 "No module named 'starlette'" ——
    真正原因（解释器不对）被埋在了模块错误里。
    """

    def test_current_env_is_complete(self) -> None:
        from saviland.interp import missing_modules
        self.assertEqual(missing_modules(), [],
                         "测试环境本身应装齐依赖，否则守卫测不出来")

    def test_finds_a_working_interpreter_or_none(self) -> None:
        from saviland.interp import find_working_python
        found = find_working_python()
        if found is not None:
            self.assertTrue(Path(found).is_file(), f"返回的解释器不存在：{found}")

    def test_site_dir_probe_sees_installed_package(self) -> None:
        """目录探测要能认出"装了"。不依赖子进程 —— 受限环境里子进程抓输出会失败。

        ★ 这里刻意不用 tempfile.TemporaryDirectory：本项目的沙箱里
        临时目录下不能建子目录（WinError 5），统一用 data/_test_tmp 下的普通目录。
        """
        from saviland.interp import _has_module

        probe = Path(__file__).resolve().parents[1] / "data" / "_test_tmp" / "interp_probe"
        probe.mkdir(parents=True, exist_ok=True)
        try:
            (probe / "starlette").mkdir(exist_ok=True)
            (probe / "singlefile.py").write_text("x = 1", encoding="utf-8")
            self.assertTrue(_has_module([probe], "starlette"))
            self.assertTrue(_has_module([probe], "singlefile"))
            self.assertFalse(_has_module([probe], "definitely_not_here"))
        finally:
            shutil.rmtree(probe, ignore_errors=True)

    def test_ensure_is_noop_when_deps_present(self) -> None:
        from saviland import interp
        interp.ensure("test")      # 依赖齐全时不应退出、不应换解释器

    def test_ensure_reports_clearly_when_no_interpreter_found(self) -> None:
        """找不到可用解释器时：退出码 2，并且给出可照抄的命令（不是抛 ImportError）。"""
        import contextlib
        import io
        from unittest import mock

        from saviland import interp
        buf = io.StringIO()
        with mock.patch.object(interp, "missing_modules", return_value=["starlette"]), \
             mock.patch.object(interp, "find_working_python", return_value=None), \
             contextlib.redirect_stdout(buf):
            with self.assertRaises(SystemExit) as ctx:
                interp.ensure("tests/run_all.py")
        self.assertEqual(ctx.exception.code, 2)
        text = buf.getvalue()
        self.assertIn("starlette", text, "要报出到底缺哪个依赖")
        self.assertIn("SAVILAND_PYTHON", text, "要给出可照抄的解决办法")


# ==========================================================================
class TestCopyExport(TempWorkspace):
    """`all_copies.csv` 的「检索命中 / 违禁词」两列必须是真的。

    ★ 回归：第一版导出把这两列硬编码成 0 和空列表，于是整张表看起来
    "RAG 没生效、合规没拦任何词" —— 一个静默的数据诚实性问题
    （表格不报错，但传达的结论是错的）。数据其实在生成时写进了侧车 JSON。
    """

    def test_sidecar_supplies_hits_and_violations(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        out = config.OUTPUTS_DIR / "listing"
        out.mkdir(parents=True, exist_ok=True)
        (out / "SV-TEST_v1.json").write_text(json.dumps({
            "copy": {"title": "t"},
            "violations": [{"term": "best", "rule": "absolute"}],
            "rag": {"mode": "hybrid", "hits": 7},
        }, ensure_ascii=False), encoding="utf-8")

        hits, viol = listing_pipe._copy_sidecar("SV-TEST", 1)
        self.assertEqual(hits, 7)
        self.assertEqual(len(viol), 1)
        self.assertEqual(viol[0]["term"], "best")

    def test_missing_or_broken_sidecar_degrades_to_zero(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        self.assertEqual(listing_pipe._copy_sidecar("SV-NOSUCH", 9), (0, []))
        broken = config.OUTPUTS_DIR / "listing" / "SV-BAD_v1.json"
        broken.parent.mkdir(parents=True, exist_ok=True)
        broken.write_text("{ not json", encoding="utf-8")
        self.assertEqual(listing_pipe._copy_sidecar("SV-BAD", 1), (0, []))

    def test_exported_row_reflects_sidecar(self) -> None:
        from saviland.pipelines import listing as listing_pipe
        seed.seed_all(force=True, days_orders=10, days_ads=10)
        # use_rag=False：测试套件必须**完全离线**，不能因为检索去调 embedding 接口
        res = listing_pipe.generate_one("SV-UVL-48W", variants=1, engine="offline",
                                        use_rag=False)

        # 把侧车改成"检索命中 5、拦了 1 个词"，再导出，CSV 必须跟着变
        side = config.OUTPUTS_DIR / "listing" / f"SV-UVL-48W_v{res.version}.json"
        data = json.loads(side.read_text(encoding="utf-8"))
        data["rag"] = {"mode": "hybrid", "hits": 5}
        data["violations"] = [{"term": "guaranteed"}]
        side.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

        csv_path = listing_pipe.export_all_latest()
        import csv as csv_mod
        with csv_path.open(encoding="utf-8-sig") as fh:
            rows = list(csv_mod.DictReader(fh))
        row = next(r for r in rows if r["sku"] == "SV-UVL-48W")
        self.assertEqual(int(row["rag_hits"]), 5, "rag_hits 必须来自侧车，不能是硬编码 0")
        self.assertIn("guaranteed", row["violations"])


# ==========================================================================
class TestAuth(unittest.TestCase):
    """后台认证：未登录不能看后台，但买家端必须公开。

    ★ 之前的状态是"前台没放后台入口"，但那**不是防护** —— 打开 8501 就是管理员。
    真实世界里后台靠"身份 + 网络 + 权限"三层，这里做最小但真实的第一层。
    """

    def setUp(self) -> None:
        self._old = os.environ.pop("SAVILAND_NO_AUTH", None)
        from starlette.testclient import TestClient

        from saviland.web.server import create_app
        self.client = TestClient(create_app(), follow_redirects=False)

    def tearDown(self) -> None:
        if self._old is not None:
            os.environ["SAVILAND_NO_AUTH"] = self._old
        else:
            os.environ.pop("SAVILAND_NO_AUTH", None)

    def _login(self, password: str = "saviland-demo", nxt: str = "/"):
        return self.client.post("/login", data={"username": "admin",
                                                "password": password, "next": nxt})

    # ---- 未登录 ----
    def test_backoffice_page_redirects_to_login(self) -> None:
        r = self.client.get("/")
        self.assertEqual(r.status_code, 303)
        self.assertIn("/login", r.headers.get("location", ""))

    def test_backoffice_api_returns_401(self) -> None:
        for path in ("/api/meta", "/api/overview", "/api/ads", "/api/runs"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 401,
                                 f"{path} 未登录必须 401，不能返回空数据装成功")

    def test_public_surface_stays_open(self) -> None:
        """买家端、静态资源、探活接口必须免登录 —— 否则买家看不了图、
        启动器"是否已在运行"的探测也会被拦。"""
        for path in ("/shop", "/shop/shop.css", "/api/health", "/login"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200, path)

    # ---- 登录 ----
    def test_wrong_password_is_rejected(self) -> None:
        self.assertEqual(self._login("wrong-password").status_code, 401)
        self.assertNotIn("sv_session", self._login("wrong-password").cookies)

    def test_login_grants_access_and_sets_httponly_cookie(self) -> None:
        r = self._login()
        self.assertEqual(r.status_code, 303)
        self.assertTrue(r.cookies.get("sv_session"))
        set_cookie = r.headers.get("set-cookie", "").lower()
        self.assertIn("httponly", set_cookie, "会话 Cookie 必须是 HttpOnly")
        self.assertIn("samesite=lax", set_cookie)
        self.assertEqual(self.client.get("/api/meta").status_code, 200)
        who = self.client.get("/api/whoami")
        self.assertEqual(who.status_code, 200)
        self.assertEqual(json.loads(who.content)["user"], "admin")

    def test_logout_invalidates_session(self) -> None:
        self._login()
        self.assertEqual(self.client.get("/api/meta").status_code, 200)
        self.client.get("/logout")
        self.assertEqual(self.client.get("/api/meta").status_code, 401)

    def test_open_redirect_is_blocked(self) -> None:
        """`next=//evil.com` 不能把人带去外站。"""
        r = self._login(nxt="//evil.example.com")
        self.assertEqual(r.headers.get("location"), "/")

    def test_no_auth_env_disables_guard(self) -> None:
        os.environ["SAVILAND_NO_AUTH"] = "1"
        self.assertEqual(self.client.get("/api/meta").status_code, 200)


# ==========================================================================
class TestStorefront(unittest.TestCase):
    """买家端：能看、且只看得到买家该看的。"""

    def setUp(self) -> None:
        os.environ["SAVILAND_NO_AUTH"] = "1"
        from starlette.testclient import TestClient

        from saviland.web.server import create_app
        self.client = TestClient(create_app())

    def tearDown(self) -> None:
        os.environ.pop("SAVILAND_NO_AUTH", None)

    def test_routes_render(self) -> None:
        for path in ("/shop", "/shop/shop.css", "/shop/p/SV-UVL-48W"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200, path)

    def test_list_shows_all_products(self) -> None:
        html = self.client.get("/shop").text
        self.assertGreaterEqual(html.count('class="card"'), 10)

    def test_detail_shows_generated_copy_and_images(self) -> None:
        html = self.client.get("/shop/p/SV-UVL-48W").text
        self.assertIn("About this item", html)
        self.assertIn("Add to Cart", html)
        self.assertGreaterEqual(html.count('class="thumb"'), 1, "应展示生成的主图")

    def test_unknown_sku_redirects_home(self) -> None:
        r = self.client.get("/shop/p/NOT-EXIST", follow_redirects=False)
        self.assertIn(r.status_code, (303, 307))
        self.assertIn("/shop", r.headers.get("location", ""))

    def test_buyer_projection_has_no_backoffice_fields(self) -> None:
        """字段是白名单投影：成本、毛利、广告、库存明细一律不出现。"""
        from saviland.web.storefront import FORBIDDEN_FIELDS, buyer_item
        item = buyer_item("SV-UVL-48W")
        self.assertIsNotNone(item)
        keys = set(item)
        for bad in FORBIDDEN_FIELDS:
            with self.subTest(field=bad):
                self.assertNotIn(bad, keys, f"买家视图不该含 {bad}")

    def test_rendered_html_leaks_no_backoffice_terms(self) -> None:
        tokens = ("cost_usd", "毛利", "gross_profit", "ACOS", "花费", "campaign",
                  "warehouse", "安全库存", "zhipu", "glm-4", "deepseek", "run_id")
        for path in ("/shop", "/shop/p/SV-UVL-48W"):
            html = self.client.get(path).text
            for t in tokens:
                with self.subTest(path=path, token=t):
                    self.assertNotIn(t, html, f"买家端 {path} 出现了后台口径：{t}")

    def test_storefront_does_not_import_backoffice(self) -> None:
        """前台模块不许 import 广告/库存/订单/运行档案等后台域模块。

        这是"前后台分开"的**可执行定义** —— 只靠目录分开不算分开，
        一旦有人图省事在 storefront 里 import 了库存，前台就会不知不觉泄露后台口径。
        """
        import ast

        src = (Path(__file__).resolve().parents[1] / "src" / "saviland" / "web"
               / "storefront.py").read_text(encoding="utf-8")
        tree = ast.parse(src)
        forbidden = {"ads", "inventory", "orders", "runs", "api", "jobs", "llm", "rag_bridge"}
        for node in ast.walk(tree):
            names: list[str] = []
            if isinstance(node, ast.Import):
                names = [a.name for a in node.names]
            elif isinstance(node, ast.ImportFrom):
                names = [node.module or ""] + [a.name for a in node.names]
            for name in names:
                tail = name.split(".")[-1]
                with self.subTest(imported=name):
                    self.assertNotIn(tail, forbidden,
                                     f"storefront 不应 import {name}（后台域模块）")

    def test_stock_label_is_coarse(self) -> None:
        """买家只该看到"有货 / 仅剩 N 件 / 缺货"，不该看到分仓与补货参数。"""
        from saviland.web.storefront import _stock_label
        self.assertEqual(_stock_label(0)[0], "Currently unavailable")
        self.assertFalse(_stock_label(0)[1])
        self.assertIn("Only 5 left", _stock_label(5)[0])
        self.assertEqual(_stock_label(500)[0], "In Stock")


# ==========================================================================
class TestPortProbe(unittest.TestCase):
    """端口探测必须能识别"别人已经在监听"。

    ★ 回归：第一版用 `SO_REUSEADDR` 做绑定测试，而 Windows 上它允许绑定一个
    已被监听的端口 —— 探测说"空闲"、uvicorn 也绑定成功，但流量被旧进程抢走，
    表现是"端口在听却一直不响应"。用户看到的就是服务起不来。
    """

    def test_free_port_is_reported_free(self) -> None:
        from saviland.web.server import _port_free
        self.assertTrue(_port_free("127.0.0.1", 8567))

    def test_listened_port_is_reported_busy(self) -> None:
        import socket

        from saviland.web.server import _port_free
        server = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        try:
            server.bind(("127.0.0.1", 8568))
            server.listen(1)
            self.assertFalse(_port_free("127.0.0.1", 8568),
                             "已被监听的端口必须判定为不可用")
        finally:
            server.close()

    def test_pick_port_shifts_when_busy(self) -> None:
        import socket

        from saviland.web.server import _pick_port
        server = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        try:
            server.bind(("127.0.0.1", 8569))
            server.listen(1)
            self.assertEqual(_pick_port("127.0.0.1", 8569), 8570)
        finally:
            server.close()


# ==========================================================================
class TestWebApi(TempWorkspace):
    """Web 层（Starlette）：接口契约 + 安全边界 + 不泄露内部实现。

    用 Starlette 的 TestClient 在进程内直接调 ASGI 应用 —— 不挑端口、不起服务，
    失败时堆栈就在当前进程里。
    """

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=21, days_ads=21)
        # 这一组测的是接口契约，不是认证（认证在 TestAuth 里单独测），
        # 所以关掉守卫，免得每个用例都要先登录一遍。
        os.environ["SAVILAND_NO_AUTH"] = "1"
        from starlette.testclient import TestClient

        from saviland.web.server import create_app
        cls.client = TestClient(create_app())

    @classmethod
    def tearDownClass(cls) -> None:
        os.environ.pop("SAVILAND_NO_AUTH", None)
        super().tearDownClass()

    def _json(self, path: str) -> dict:
        r = self.client.get(path)
        self.assertEqual(r.status_code, 200, f"{path} → {r.status_code}: {r.text[:200]}")
        return json.loads(r.content.decode("utf-8"))

    def test_index_and_static(self) -> None:
        r = self.client.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Saviland", r.text)
        for path, needle in (("/static/app.js", "Saviland"),
                             ("/static/style.css", "--brand")):
            r = self.client.get(path)
            self.assertEqual(r.status_code, 200, path)
            self.assertIn(needle, r.text, path)

    def test_all_json_endpoints(self) -> None:
        cases = {
            "/api/health": ["ok"],
            "/api/meta": ["brand", "counts", "knowledge"],
            "/api/overview": ["kpis", "week"],
            "/api/products": ["items", "categories"],
            "/api/copies": ["items"],
            "/api/inventory": ["items", "levels", "thresholds"],
            "/api/orders?days=30": ["summary", "by_channel", "daily"],
            "/api/ads": ["total", "by_campaign", "daily"],
            "/api/findings": ["items", "summary", "listing_targets"],
            "/api/report": ["exists", "downloads", "charts"],
            "/api/runs": ["items", "engine_usage"],
            "/api/rag/status": ["chunks", "docs"],
        }
        for path, keys in cases.items():
            with self.subTest(path=path):
                data = self._json(path)
                for k in keys:
                    self.assertIn(k, data, f"{path} 缺少字段 {k}")

    def test_responses_are_json_safe(self) -> None:
        """回归：`available_days` 在没有动销时是 float('inf')，
        直接序列化会产出 `Infinity`，浏览器 JSON.parse 会直接抛错。"""
        raw = self.client.get("/api/inventory").content.decode("utf-8")
        for token in ("Infinity", "NaN"):
            self.assertNotIn(token, raw, f"响应里不允许出现 {token}")

    def test_path_traversal_is_rejected(self) -> None:
        for path in ("/media/product/../../../data/saviland.db",
                     "/media/product/....//....//README.md",
                     "/media/nosuch/x.png"):
            with self.subTest(path=path):
                r = self.client.get(path)
                self.assertIn(r.status_code, (400, 403, 404), f"{path} 未被拒绝")

    def test_product_detail_and_404(self) -> None:
        data = self._json("/api/products/SV-UVL-48W")
        self.assertIn("specs", data)
        self.assertIn("hero_images", data)
        self.assertEqual(self.client.get("/api/products/NOT-EXIST").status_code, 404)

    def test_validation_errors(self) -> None:
        self.assertEqual(self.client.get("/api/rag/search").status_code, 400)
        r = self.client.post("/api/generate/listing", json={"skus": []})
        self.assertEqual(r.status_code, 400)

    def test_job_lifecycle(self) -> None:
        """POST 起任务应立刻返回 job_id，且能被轮询到。

        ★ 必须传 `engine: "offline"`：否则这个"只验证任务生命周期"的用例会真的
        调用一次付费模型，把整个测试套件从 46 秒拖到 2 分钟，还烧额度。
        测试永远不该访问网络。
        """
        r = self.client.post("/api/generate/listing",
                             json={"skus": ["SV-UVL-48W"], "variants": 1, "engine": "offline"})
        self.assertEqual(r.status_code, 200)
        job_id = json.loads(r.content.decode("utf-8"))["job_id"]
        job = self._json(f"/api/jobs/{job_id}")
        self.assertIn(job["status"], ("queued", "running", "done", "failed"))
        self.assertIn("log", job)
        self.assertNotEqual(self.client.get("/api/jobs/nope").status_code, 200)

    def test_unknown_engine_is_ignored(self) -> None:
        r = self.client.post("/api/generate/listing",
                             json={"skus": ["SV-GEL-BASE"], "variants": 1,
                                   "engine": "../../etc/passwd"})
        self.assertEqual(r.status_code, 200, "非法 engine 应被忽略而不是丢 500")

    # ---- 面向用户的信息面：不许出现实现细节 ----
    def test_meta_does_not_leak_engine(self) -> None:
        raw = self.client.get("/api/meta").content.decode("utf-8").lower()
        for word in ("zhipu", "glm", "deepseek", "ollama", "embedding"):
            self.assertNotIn(word, raw, f"/api/meta 不应出现内部实现词：{word}")

    def test_findings_do_not_leak_engine(self) -> None:
        raw = self.client.get("/api/findings").content.decode("utf-8").lower()
        for word in ("zhipu", "glm-4", "deepseek", "prompt"):
            self.assertNotIn(word, raw, f"/api/findings 不应出现内部实现词：{word}")

    def test_runs_view_keeps_engine_for_audit(self) -> None:
        """运行档案是技术视图，引擎信息**只允许**出现在这里。"""
        data = self._json("/api/runs")
        self.assertIsInstance(data["engine_usage"], list)


# ==========================================================================
class TestReportingNoEngine(TempWorkspace):
    """周报正脸不许出现引擎等内部信息。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        seed.seed_all(force=True, days_orders=21, days_ads=21)
        from saviland.pipelines import weekly as weekly_pipe
        cls.res = weekly_pipe.generate_weekly(engine="offline", use_rag=False)
        cls.md = Path(cls.res.report_md).read_text(encoding="utf-8")
        cls.html = Path(cls.res.report_html).read_text(encoding="utf-8")

    def test_markdown_has_no_engine_or_latency(self) -> None:
        for word in ("生成引擎", "耗时", "prompt", "zhipu", "glm-4", "deepseek"):
            with self.subTest(word=word):
                self.assertNotIn(word, self.md, f"周报 Markdown 不应出现：{word}")

    def test_html_has_no_engine(self) -> None:
        for word in ("zhipu", "glm-4", "deepseek", "生成引擎"):
            with self.subTest(word=word):
                self.assertNotIn(word, self.html.lower(), f"周报 HTML 不应出现：{word}")

    def test_report_keeps_period_and_timestamp(self) -> None:
        self.assertIn("报告周期", self.md)
        self.assertIn("生成时间", self.md)

    def test_report_keeps_traceability_without_plumbing(self) -> None:
        """报告编号可以留（对账要用），但引擎/耗时不行。"""
        self.assertIn("报告编号", self.md)
        self.assertIn("数据窗口", self.md)

    def test_html_is_self_contained_and_styled(self) -> None:
        self.assertIn("<!DOCTYPE html>", self.html)
        self.assertIn("data:image/png;base64,", self.html, "图表必须内嵌")
        self.assertIn("@media print", self.html, "应自带打印样式")
        self.assertIn("class=\"hero\"", self.html, "应有品牌头部")


# ==========================================================================
class TestRunAudit(TempWorkspace):
    """运行档案的完整性：失败与中断都必须留痕，不许留下永不结束的 running。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        if not (db.scalar("SELECT COUNT(*) FROM products") or 0):
            seed.seed_all(force=True, days_orders=14, days_ads=14)

    def test_failed_generation_is_recorded_as_failed(self) -> None:
        """回归：批量生成会逐个 SKU 捕获异常，若异常发生在 start_run 之后，
        那条运行记录会**永远停在 running** —— 批量任务看起来正常完成，
        档案里却躺着僵尸记录。"""
        from saviland import runs
        from saviland.pipelines import listing as listing_pipe

        original = catalog_domain.save_listing_copy

        def boom(*_a, **_k):
            raise RuntimeError("模拟数据库写入故障")

        catalog_domain.save_listing_copy = boom
        try:
            out = listing_pipe.generate_batch(["SV-UVL-48W"], variants=1,
                                              engine="offline", use_rag=False)
        finally:
            catalog_domain.save_listing_copy = original

        self.assertEqual(out["failed"], 1)
        row = db.query_one(
            "SELECT status, note FROM generation_runs WHERE pipeline = 'listing' "
            "ORDER BY rowid DESC LIMIT 1")
        self.assertEqual(row["status"], "failed")
        self.assertIn("模拟数据库写入故障", row["note"])
        _ = runs

    def test_sweep_marks_stale_running_as_interrupted(self) -> None:
        from saviland import runs
        run_id = runs.start_run("weekly", target="sweep-test")
        db.execute("UPDATE generation_runs SET created_at = datetime('now', '-2 hours') "
                   "WHERE run_id = ?", (run_id,))
        swept = runs.sweep_stale(minutes=30)
        self.assertGreaterEqual(swept, 1)
        row = db.query_one("SELECT status FROM generation_runs WHERE run_id = ?", (run_id,))
        self.assertEqual(row["status"], "interrupted")

    def test_fresh_running_is_not_swept(self) -> None:
        from saviland import runs
        run_id = runs.start_run("weekly", target="fresh-test")
        runs.sweep_stale(minutes=30)
        row = db.query_one("SELECT status FROM generation_runs WHERE run_id = ?", (run_id,))
        self.assertEqual(row["status"], "running", "刚开跑的任务不该被误判为中断")


# ==========================================================================
if __name__ == "__main__":
    unittest.main(verbosity=2)
